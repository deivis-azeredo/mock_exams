"""
Núcleo de lógica do gerador de provas — versão web.

Mesma lógica do script desktop original, mas:
- Lê o PDF de um objeto em memória (upload do navegador), não de um caminho de disco.
- Gera o .docx em memória (bytes) para download pelo navegador, não salva em pasta.
- Não importa tkinter nem groq/dotenv no topo do arquivo — o cliente da IA é
  passado como parâmetro, construído pela camada web com a chave vinda dos
  "Secrets" do Streamlit.
"""

import io
import json
import re
import datetime

from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Cores usadas no documento
COR_TITULO = RGBColor(24, 43, 73)
COR_QUESTAO = RGBColor(30, 60, 114)
COR_CORRETA = RGBColor(0, 128, 0)

# O modelo tem uma janela de contexto limitada (entrada + saída somadas).
LIMITE_CARACTERES_TEXTO = 45000


# ---------------------------------------------------------------------------
# 1. Extração de texto do PDF (a partir de um arquivo em memória)
# ---------------------------------------------------------------------------
def extrair_texto(arquivo_pdf):
    """arquivo_pdf: objeto tipo arquivo (ex: o que vem de st.file_uploader)."""
    leitor = PdfReader(arquivo_pdf)
    texto_completo = ""
    for pagina in leitor.pages:
        conteudo = pagina.extract_text()
        if conteudo:
            texto_completo += conteudo + "\n"
    return texto_completo.strip()


def preparar_texto_para_ia(texto_completo, limite_caracteres=LIMITE_CARACTERES_TEXTO):
    """Amostra início + meio + fim de textos grandes, para caber no contexto do modelo."""
    if len(texto_completo) <= limite_caracteres:
        return texto_completo, False

    numero_blocos = 6
    tamanho_bloco = limite_caracteres // numero_blocos
    tamanho_total = len(texto_completo)

    partes = []
    for i in range(numero_blocos):
        inicio = int((i / numero_blocos) * tamanho_total)
        fim = min(inicio + tamanho_bloco, tamanho_total)
        if i == numero_blocos - 1:
            fim = tamanho_total
            inicio = max(0, fim - tamanho_bloco)
        partes.append(texto_completo[inicio:fim])

    return "\n\n[...trecho omitido...]\n\n".join(partes), True


# ---------------------------------------------------------------------------
# 2. Chamada da IA
# ---------------------------------------------------------------------------
def gerar_prova_ia(client, texto_base, nome_arquivo, quantidade_questoes):
    """client: instância já criada de groq.Groq(api_key=...)."""
    prompt = f"""
Você é um professor e elaborador de avaliações experiente.
Com base no texto do documento "{nome_arquivo}" fornecido abaixo, crie uma avaliação pedagógica completa.

A avaliação deve conter:
- Exatamente {quantidade_questoes} questões de múltipla escolha, cada uma com 4 alternativas (A, B, C, D) e apenas uma correta.
- Exatamente 2 questões dissertativas.
- Para cada questão de múltipla escolha, indique a letra correta e uma justificativa curta (1 frase).
- Para cada dissertativa, indique um resumo do que se espera na resposta (1 a 2 frases).

Responda SOMENTE com um JSON válido, sem comentários, sem markdown, sem texto antes ou depois, seguindo EXATAMENTE este formato:

{{
  "questoes_multipla_escolha": [
    {{
      "enunciado": "texto da pergunta",
      "alternativas": {{"A": "texto", "B": "texto", "C": "texto", "D": "texto"}},
      "resposta_correta": "B",
      "justificativa": "explicação curta"
    }}
  ],
  "questoes_dissertativas": [
    {{
      "enunciado": "texto da pergunta dissertativa",
      "resposta_esperada": "resumo do que se espera na resposta"
    }}
  ]
}}

TEXTO BASE:
---
{texto_base}
---
"""

    max_tokens_resposta = min(4000 + quantidade_questoes * 220, 30000)

    resposta = client.chat.completions.create(
        messages=[
            {"role": "system", "content": "Você é um assistente educacional que responde apenas com JSON válido, sem nenhum texto adicional."},
            {"role": "user", "content": prompt},
        ],
        model="openai/gpt-oss-120b",
        temperature=0.4,
        max_tokens=max_tokens_resposta,
    )
    return resposta.choices[0].message.content


def extrair_json(texto_resposta):
    texto = texto_resposta.strip()
    texto = re.sub(r"^```json\s*|^```\s*|```$", "", texto, flags=re.MULTILINE).strip()
    try:
        return json.loads(texto)
    except json.JSONDecodeError:
        pass
    inicio = texto.find("{")
    fim = texto.rfind("}")
    if inicio != -1 and fim != -1 and fim > inicio:
        return json.loads(texto[inicio:fim + 1])
    raise ValueError("Não foi possível interpretar o JSON retornado pela IA.")


# ---------------------------------------------------------------------------
# 3. Montagem do .docx (em memória)
# ---------------------------------------------------------------------------
def _definir_bordas_paragrafo(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom_el = OxmlElement('w:bottom')
    bottom_el.set(qn('w:val'), 'single')
    bottom_el.set(qn('w:sz'), '6')
    bottom_el.set(qn('w:space'), '1')
    bottom_el.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom_el)
    pPr.append(pBdr)


def _adicionar_cabecalho_identificacao(doc):
    tabela = doc.add_table(rows=2, cols=4)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.style = 'Table Grid'
    cabecalhos = ["Nome do Aluno", "Data", "Turma", "Nota"]
    larguras = [Inches(3.2), Inches(1.2), Inches(1.0), Inches(0.9)]
    for i, texto in enumerate(cabecalhos):
        celula = tabela.rows[0].cells[i]
        p = celula.paragraphs[0]
        run = p.add_run(texto)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Arial'
        celula.width = larguras[i]
    for i in range(4):
        celula = tabela.rows[1].cells[i]
        celula.paragraphs[0].add_run(" ")
        celula.width = larguras[i]
    tabela.rows[1].height = Inches(0.35)
    doc.add_paragraph()


def _adicionar_questao_multipla_escolha(doc, numero, questao):
    p_enunciado = doc.add_paragraph()
    run_num = p_enunciado.add_run(f"{numero}. ")
    run_num.bold = True
    run_num.font.name = 'Arial'
    run_num.font.size = Pt(11)
    run_num.font.color.rgb = COR_QUESTAO
    run_texto = p_enunciado.add_run(questao.get("enunciado", ""))
    run_texto.font.name = 'Arial'
    run_texto.font.size = Pt(11)
    p_enunciado.paragraph_format.space_after = Pt(4)

    alternativas = questao.get("alternativas", {}) or {}
    for letra in ["A", "B", "C", "D"]:
        texto_alt = alternativas.get(letra, "")
        if not texto_alt:
            continue
        p_alt = doc.add_paragraph()
        p_alt.paragraph_format.left_indent = Inches(0.4)
        p_alt.paragraph_format.space_after = Pt(2)
        run_letra = p_alt.add_run(f"{letra}) ")
        run_letra.bold = True
        run_letra.font.name = 'Arial'
        run_letra.font.size = Pt(10.5)
        run_texto_alt = p_alt.add_run(texto_alt)
        run_texto_alt.font.name = 'Arial'
        run_texto_alt.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def _adicionar_questao_dissertativa(doc, numero, questao):
    p_enunciado = doc.add_paragraph()
    run_num = p_enunciado.add_run(f"{numero}. ")
    run_num.bold = True
    run_num.font.name = 'Arial'
    run_num.font.size = Pt(11)
    run_num.font.color.rgb = COR_QUESTAO
    run_texto = p_enunciado.add_run(questao.get("enunciado", ""))
    run_texto.font.name = 'Arial'
    run_texto.font.size = Pt(11)
    p_enunciado.paragraph_format.space_after = Pt(6)

    for _ in range(4):
        linha = doc.add_paragraph()
        linha.paragraph_format.space_after = Pt(10)
        _definir_bordas_paragrafo(linha)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def _adicionar_gabarito(doc, questoes_objetivas, questoes_dissertativas):
    doc.add_page_break()
    titulo = doc.add_heading("GABARITO", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if titulo.runs:
        titulo.runs[0].font.name = 'Arial'
        titulo.runs[0].font.size = Pt(16)
        titulo.runs[0].font.color.rgb = COR_TITULO
    doc.add_paragraph()

    subtitulo = doc.add_paragraph()
    run = subtitulo.add_run("Questões de Múltipla Escolha")
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = COR_QUESTAO
    subtitulo.paragraph_format.space_after = Pt(6)

    for i, questao in enumerate(questoes_objetivas, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run_num = p.add_run(f"{i}. ")
        run_num.bold = True
        run_num.font.name = 'Arial'
        run_num.font.size = Pt(10.5)
        run_resposta = p.add_run(f"Resposta: {questao.get('resposta_correta', '-')}  ")
        run_resposta.bold = True
        run_resposta.font.name = 'Arial'
        run_resposta.font.size = Pt(10.5)
        run_resposta.font.color.rgb = COR_CORRETA
        justificativa = questao.get("justificativa", "")
        if justificativa:
            run_just = p.add_run(f"— {justificativa}")
            run_just.italic = True
            run_just.font.name = 'Arial'
            run_just.font.size = Pt(10)

    if questoes_dissertativas:
        doc.add_paragraph()
        subtitulo2 = doc.add_paragraph()
        run2 = subtitulo2.add_run("Questões Dissertativas — Respostas Esperadas")
        run2.bold = True
        run2.font.name = 'Arial'
        run2.font.size = Pt(12)
        run2.font.color.rgb = COR_QUESTAO
        subtitulo2.paragraph_format.space_after = Pt(6)
        for i, questao in enumerate(questoes_dissertativas, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run_num = p.add_run(f"{i}. ")
            run_num.bold = True
            run_num.font.name = 'Arial'
            run_num.font.size = Pt(10.5)
            run_texto = p.add_run(questao.get("resposta_esperada", ""))
            run_texto.font.name = 'Arial'
            run_texto.font.size = Pt(10.5)


def construir_docx_bytes(dados, nome_origem, aviso_teste=None):
    """
    Monta o documento e devolve os BYTES do .docx (para st.download_button),
    em vez de salvar em disco.

    aviso_teste: se fornecido (string), adiciona uma linha discreta abaixo do
    título — usado pela versão de teste do site para marcar o documento.
    """
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    titulo = doc.add_heading(f"AVALIAÇÃO - {nome_origem.upper()}", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if titulo.runs:
        titulo.runs[0].font.name = 'Arial'
        titulo.runs[0].font.size = Pt(16)
        titulo.runs[0].font.color.rgb = COR_TITULO

    if aviso_teste:
        aviso = doc.add_paragraph()
        aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_aviso = aviso.add_run(aviso_teste)
        run_aviso.italic = True
        run_aviso.font.size = Pt(8)
        run_aviso.font.color.rgb = RGBColor(150, 150, 150)
        run_aviso.font.name = 'Arial'

    doc.add_paragraph()
    _adicionar_cabecalho_identificacao(doc)

    questoes_objetivas = dados.get("questoes_multipla_escolha", []) or []
    questoes_dissertativas = dados.get("questoes_dissertativas", []) or []

    for i, questao in enumerate(questoes_objetivas, 1):
        _adicionar_questao_multipla_escolha(doc, i, questao)

    if questoes_dissertativas:
        subtitulo = doc.add_paragraph()
        run = subtitulo.add_run("Questões Dissertativas")
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.color.rgb = COR_QUESTAO
        subtitulo.paragraph_format.space_before = Pt(6)
        subtitulo.paragraph_format.space_after = Pt(8)
        for i, questao in enumerate(questoes_dissertativas, 1):
            _adicionar_questao_dissertativa(doc, i, questao)

    _adicionar_gabarito(doc, questoes_objetivas, questoes_dissertativas)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
